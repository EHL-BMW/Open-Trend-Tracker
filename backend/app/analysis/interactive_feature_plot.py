"""Generate interactive feature-impact and BMW sales forecast plots."""
from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import load_workbook

from app.analysis.official_sales import (
    feature_impact_forecast,
    read_versions,
    version_sales_differences,
)
from app.paths import DATA_DIR, EXTERNAL_DIR

DEFAULT_INPUT = EXTERNAL_DIR / "chinese_official_200_pairs.csv"
DEFAULT_OUTPUT = DATA_DIR / "analysis/chinese_200/plots/interactive_feature_impact.html"
DEFAULT_BMW_SALES = DATA_DIR / "internal/sales_volumes.xlsx"

CATEGORY_COLORS = {
    "Energy & Charging": "#2f78a7",
    "Driving Assistance": "#a353a3",
    "Digital Cockpit & AI": "#e08332",
    "Vehicle Dynamics": "#397c5a",
    "Smart Cabin": "#b85c68",
    "Connected Ecosystem": "#5e6fc2",
    "Smart Exterior": "#8a733e",
    "General & Services": "#66717d",
    "Unclassified": "#7b7b7b",
}


def build_snapshots(input_path: Path) -> dict[int, list[dict]]:
    """Build year-cutoff snapshots; the chart uses the newest complete one."""
    comparisons = version_sales_differences(read_versions(input_path))
    years = sorted({int(str(row["release_date"])[:4]) for row in comparisons})
    snapshots: dict[int, list[dict]] = {}
    for cutoff in years:
        eligible = [row for row in comparisons if int(str(row["release_date"])[:4]) <= cutoff]
        forecast = feature_impact_forecast(eligible)
        snapshots[cutoff] = [
            {
                "feature": str(row["display_name"]),
                "category": str(row["category"]),
                "pairs": int(row["pair_count"]),
                "models": int(row["model_count"]),
                "estimate": [round(float(row[f"estimated_sales_impact_year_{year}_pct"]), 3)
                             for year in range(1, 5)],
                "lower": [round(float(row[f"lower_year_{year}_pct"]), 3) for year in range(1, 5)],
                "upper": [round(float(row[f"upper_year_{year}_pct"]), 3) for year in range(1, 5)],
            }
            for row in forecast
        ]
    return snapshots


def load_bmw_sales_history(path: Path) -> dict[str, dict[str, dict[str, float | str]]]:
    """Read BMW annual sales history and add a Global total for each model."""
    workbook = load_workbook(path, data_only=True, read_only=True)
    result: dict[str, dict[str, dict[str, float | str]]] = {}
    for worksheet in workbook.worksheets:
        rows = list(worksheet.iter_rows(values_only=True))
        if not rows:
            continue
        headers = [str(value).strip() if value is not None else "" for value in rows[0]]
        index = {name: position for position, name in enumerate(headers)}
        required = {"market", "market_code", "volume_2024", "volume_2025", "volume_2030"}
        if not required.issubset(index):
            continue
        model_markets: dict[str, dict[str, float | str]] = {}
        for row in rows[1:]:
            market_code = row[index["market_code"]]
            if not market_code:
                continue
            model_markets[str(market_code)] = {
                "label": str(row[index["market"]]),
                "2024": float(row[index["volume_2024"]] or 0),
                "2025": float(row[index["volume_2025"]] or 0),
                "2030": float(row[index["volume_2030"]] or 0),
            }
        model_markets["Global"] = {
            "label": "Global",
            "2024": sum(float(values["2024"]) for values in model_markets.values()),
            "2025": sum(float(values["2025"]) for values in model_markets.values()),
            "2030": sum(float(values["2030"]) for values in model_markets.values()),
        }
        result[worksheet.title.replace("_", "/")] = model_markets
    workbook.close()
    if not result:
        raise ValueError(f"No BMW sales history found in {path}")
    return result


def estimate_bmw_sales(
    history: dict[str, dict[str, dict[str, float | str]]],
    model: str,
    market: str,
    selected_features: list[dict],
) -> dict[str, list[float] | float]:
    """Estimate 2030-2033 sales using the history CAGR and additive feature impacts."""
    source = history[model][market]
    volume_2025 = float(source["2025"])
    volume_2030 = float(source["2030"])
    cagr = (volume_2030 / volume_2025) ** (1 / 5) - 1 if volume_2025 > 0 and volume_2030 > 0 else 0.0
    baseline = [volume_2030 * (1 + cagr) ** offset for offset in range(4)]

    def combined(key: str) -> list[float]:
        return [sum(float(feature[key][year]) for feature in selected_features) for year in range(4)]

    impact = combined("estimate")
    lower_impact = combined("lower")
    upper_impact = combined("upper")
    return {
        "years": [2030, 2031, 2032, 2033],
        "baseline": baseline,
        "estimate": [value * (1 + pct / 100) for value, pct in zip(baseline, impact)],
        "lower": [value * (1 + pct / 100) for value, pct in zip(baseline, lower_impact)],
        "upper": [value * (1 + pct / 100) for value, pct in zip(baseline, upper_impact)],
        "impact": impact,
        "cagr": cagr,
    }


def render_html(
    snapshots: dict[int, list[dict]],
    output: Path,
    input_pair_count: int | None = None,
    input_row_count: int | None = None,
    bmw_sales_history: dict[str, dict[str, dict[str, float | str]]] | None = None,
) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    latest_year = max(snapshots)
    features = snapshots[latest_year]
    payload = json.dumps(features, ensure_ascii=False, separators=(",", ":"))
    colors = json.dumps(CATEGORY_COLORS, ensure_ascii=False, separators=(",", ":"))
    bmw_sales = json.dumps(bmw_sales_history or {}, ensure_ascii=False, separators=(",", ":"))
    pair_count = input_pair_count if input_pair_count is not None else 0
    row_count = input_row_count if input_row_count is not None else pair_count * 2
    template = r'''<!doctype html>
<html lang="en">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width,initial-scale=1">
  <title>Interactive feature-impact forecast</title>
  <script src="https://cdn.jsdelivr.net/npm/plotly.js-dist-min@2.35.2/plotly.min.js"></script>
  <style>
    :root { color-scheme: light dark; --bg: light-dark(#ffffff,#15171a); --fg: light-dark(#17202a,#eef1f4); --muted: light-dark(#596777,#aeb8c4); --line: light-dark(#d9dfe6,#39414a); }
    * { box-sizing: border-box; }
    body { margin: 0; padding: 20px; background: var(--bg); color: var(--fg); font: 14px/1.4 system-ui,-apple-system,BlinkMacSystemFont,"Segoe UI",sans-serif; }
    main { max-width: 1240px; margin: 0 auto; }
    h1 { margin: 0 0 4px; font-size: 24px; font-weight: 500; }
    h2 { margin: 30px 0 4px; font-size: 20px; font-weight: 500; }
    .subtitle,.note,.status { color: var(--muted); }
    .controls { display: flex; flex-wrap: wrap; gap: 12px; margin: 16px 0 4px; }
    .field { display: grid; gap: 4px; min-width: 180px; }
    .field label { font-weight: 500; }
    select { border: 1px solid var(--line); background: var(--bg); color: var(--fg); padding: 8px 10px; font: inherit; }
    .switches { margin: 16px 0 8px; }
    .switches > label { display: block; font-weight: 500; margin-bottom: 5px; }
    .switch-row { display: flex; flex-wrap: wrap; gap: 5px 14px; }
    .switch { appearance: none; border: 0; background: transparent; color: var(--fg); padding: 5px 0; font: inherit; display: inline-flex; align-items: center; gap: 7px; cursor: pointer; }
    .switch::before { content: ""; width: 13px; height: 13px; border-radius: 3px; background: var(--cat); opacity: 1; flex: 0 0 auto; }
    .switch[aria-pressed=false] { color: var(--muted); text-decoration: line-through; }
    .switch[aria-pressed=false]::before { opacity: .2; }
    .status { margin: 8px 0; min-height: 20px; }
    .plot { width: 100%; min-height: 520px; border: 1px solid var(--line); margin-top: 14px; }
    .note { margin-top: 10px; font-size: 12px; }
    @media (max-width: 650px) { body { padding: 12px; } .plot { min-height: 470px; } }
  </style>
</head>
<body>
<main>
  <h1>Four-year cumulative-sales improvement scenario</h1>
  <div class="subtitle">Calculated from all __PAIR_COUNT__ input pairs (__ROW_COUNT__ version rows). Toggle categories and concrete feature implementations; both plots show uncertainty bands.</div>
  <div class="switches"><label>Categories</label><div id="category-switches" class="switch-row"></div></div>
  <div class="switches"><label>Feature implementations</label><div id="feature-switches" class="switch-row"></div></div>
  <div id="status" class="status" aria-live="polite"></div>
  <div id="categories" class="plot" role="img" aria-label="Category cumulative-sales improvement over years one through four with uncertainty bands"></div>
  <div id="features" class="plot" role="img" aria-label="Feature cumulative-sales improvement over years one through four with uncertainty bands"></div>
  <div class="note">Planning scenario only: category priors × feature multiplier × adoption saturation. It is not fitted causal evidence because version-level cumulative sales are unavailable.</div>

  <h2>New BMW model sales estimate: 2030–2033</h2>
  <div class="subtitle">Choose a BMW history series and one or more concrete implementations. The no-feature baseline extends its 2025–2030 CAGR; selected implementation impacts are added at the same estimated rates shown above.</div>
  <div class="controls">
    <div class="field"><label for="bmw-model">BMW model history</label><select id="bmw-model"></select></div>
    <div class="field"><label for="bmw-market">Market</label><select id="bmw-market"></select></div>
  </div>
  <div class="switches"><label>Implementations included in the BMW estimate</label><div id="bmw-feature-switches" class="switch-row"></div></div>
  <div id="bmw-status" class="status" aria-live="polite"></div>
  <div id="bmw-sales" class="plot" role="img" aria-label="BMW sales estimate from 2030 through 2033 with selectable feature impacts and uncertainty band"></div>
  <div class="note">Scenario estimate, not a causal prediction. Impacts are added across selected features, so broad bundles can overstate gains when features overlap.</div>
</main>
<script>
const rows = __FEATURES__;
const categoryColors = __COLORS__;
const bmwSales = __BMW_SALES__;
const inputPairCount = __PAIR_COUNT__;
const horizon = [1,2,3,4];
const categories = [...new Set(rows.map(row => row.category))];
const visibleCategories = new Set(categories);
const visibleFeatures = new Set(rows.map(row => row.feature));
const categoryBox = document.getElementById('category-switches');
const featureBox = document.getElementById('feature-switches');
const bmwFeatureBox = document.getElementById('bmw-feature-switches');
const status = document.getElementById('status');
const bmwStatus = document.getElementById('bmw-status');
const bmwModel = document.getElementById('bmw-model');
const bmwMarket = document.getElementById('bmw-market');
const selectedBmwFeatures = new Set();
const dashes = ['solid','dash','dot','dashdot','longdash','longdashdot'];

function addSwitch(box,label,color,onToggle,initial=true) {
  const button=document.createElement('button');
  button.type='button'; button.className='switch'; button.textContent=label;
  button.style.setProperty('--cat',color); button.setAttribute('aria-pressed',String(initial));
  button.addEventListener('click',()=>{onToggle(button.getAttribute('aria-pressed')==='true'); render();});
  box.appendChild(button); return button;
}
for (const category of categories) addSwitch(categoryBox,category,categoryColors[category],pressed=>{
  if (pressed) visibleCategories.delete(category); else visibleCategories.add(category);
  [...categoryBox.children].find(button=>button.textContent===category).setAttribute('aria-pressed',String(!pressed));
});
for (const row of rows) addSwitch(featureBox,row.feature,categoryColors[row.category],pressed=>{
  if (pressed) visibleFeatures.delete(row.feature); else visibleFeatures.add(row.feature);
  [...featureBox.children].find(button=>button.textContent===row.feature).setAttribute('aria-pressed',String(!pressed));
});
for (const row of rows) addSwitch(bmwFeatureBox,row.feature,categoryColors[row.category],pressed=>{
  if (pressed) selectedBmwFeatures.delete(row.feature); else selectedBmwFeatures.add(row.feature);
  [...bmwFeatureBox.children].find(button=>button.textContent===row.feature).setAttribute('aria-pressed',String(!pressed));
},false);

for (const model of Object.keys(bmwSales)) bmwModel.add(new Option(model,model));
function updateMarkets() {
  const previous=bmwMarket.value;
  bmwMarket.replaceChildren();
  const markets=bmwSales[bmwModel.value] || {};
  for (const code of Object.keys(markets)) bmwMarket.add(new Option(`${markets[code].label} (${code})`,code));
  bmwMarket.value=Object.hasOwn(markets,previous)?previous:(Object.hasOwn(markets,'Global')?'Global':Object.keys(markets)[0]);
}
if (bmwModel.options.length) updateMarkets();

function theme() {
  const css=getComputedStyle(document.documentElement);
  return {bg:css.getPropertyValue('--bg').trim(),fg:css.getPropertyValue('--fg').trim(),muted:css.getPropertyValue('--muted').trim(),line:css.getPropertyValue('--line').trim()};
}
function rgba(hex,alpha) {
  const value=parseInt(hex.slice(1),16); return `rgba(${value>>16},${(value>>8)&255},${value&255},${alpha})`;
}
function bandTraces(name,estimate,lower,upper,color,dash='solid',meta='') {
  return [
    {x:horizon,y:upper,type:'scatter',mode:'lines',line:{width:0},hoverinfo:'skip',showlegend:false,legendgroup:name},
    {x:horizon,y:lower,type:'scatter',mode:'lines',fill:'tonexty',fillcolor:rgba(color,.14),line:{width:0},hoverinfo:'skip',showlegend:false,legendgroup:name},
    {x:horizon,y:estimate,type:'scatter',mode:'lines+markers',name,line:{color,width:2.5,dash},marker:{color,size:7},customdata:horizon.map(()=>meta),hovertemplate:`<b>${name}</b><br>Year %{x}: %{y:.1f}%<br>${meta}<extra></extra>`,legendgroup:name}
  ];
}
function categoryRows() {
  return categories.map(category=>{
    const members=rows.filter(row=>row.category===category);
    const weight=members.reduce((sum,row)=>sum+row.pairs,0);
    const mean=key=>horizon.map((_,i)=>members.reduce((sum,row)=>sum+row[key][i]*row.pairs,0)/weight);
    return {category,estimate:mean('estimate'),lower:mean('lower'),upper:mean('upper'),features:members.length,pairs:Math.max(...members.map(row=>row.pairs))};
  });
}
function layout(title,colors) {
  return {
    title:{text:title,x:.02,xanchor:'left'},paper_bgcolor:colors.bg,plot_bgcolor:colors.bg,font:{color:colors.fg},
    margin:{l:72,r:24,t:60,b:65},hovermode:'x unified',
    xaxis:{title:'Years after release',tickmode:'array',tickvals:horizon,range:[.85,4.15],gridcolor:colors.line,zeroline:false},
    yaxis:{title:'Cumulative-sales improvement (%)',rangemode:'tozero',gridcolor:colors.line,zeroline:false},
    legend:{orientation:'h',y:-.2},hoverlabel:{bgcolor:colors.bg,font:{color:colors.fg}},transition:{duration:220}
  };
}
function renderBmw(colors) {
  const model=bmwModel.value;
  const market=bmwMarket.value;
  const source=bmwSales[model]?.[market];
  if (!source) {
    Plotly.react('bmw-sales',[],layout('BMW sales history unavailable',colors),{responsive:true,displaylogo:false});
    bmwStatus.textContent='No BMW sales history is available.';
    return;
  }
  const years=[2030,2031,2032,2033];
  const cagr=source['2025']>0&&source['2030']>0?Math.pow(source['2030']/source['2025'],1/5)-1:0;
  const baseline=years.map((_,index)=>source['2030']*Math.pow(1+cagr,index));
  const selected=rows.filter(row=>selectedBmwFeatures.has(row.feature));
  const combined=key=>years.map((_,index)=>selected.reduce((sum,row)=>sum+row[key][index],0));
  const impact=combined('estimate');
  const lowerImpact=combined('lower');
  const upperImpact=combined('upper');
  const estimate=baseline.map((value,index)=>value*(1+impact[index]/100));
  const lower=baseline.map((value,index)=>value*(1+lowerImpact[index]/100));
  const upper=baseline.map((value,index)=>value*(1+upperImpact[index]/100));
  const featureLabel=selected.length?selected.map(row=>row.feature).join(', '):'No feature uplift';
  const traces=[
    {x:[2024,2025,2030],y:[source['2024'],source['2025'],source['2030']],type:'scatter',mode:'lines+markers',name:'BMW history / 2030 plan',line:{color:colors.muted,width:2},marker:{color:colors.muted,size:7},hovertemplate:'%{x}: %{y:,.0f} vehicles<extra>BMW history / plan</extra>'},
    {x:years,y:baseline,type:'scatter',mode:'lines+markers',name:'No-feature baseline',line:{color:colors.muted,width:2,dash:'dash'},marker:{color:colors.muted,size:7},hovertemplate:'%{x}: %{y:,.0f} vehicles<extra>No-feature baseline</extra>'},
    {x:years,y:upper,type:'scatter',mode:'lines',line:{width:0},hoverinfo:'skip',showlegend:false,legendgroup:'selected'},
    {x:years,y:lower,type:'scatter',mode:'lines',fill:'tonexty',fillcolor:rgba(categoryColors['Driving Assistance'],.16),line:{width:0},hoverinfo:'skip',showlegend:false,legendgroup:'selected'},
    {x:years,y:estimate,type:'scatter',mode:'lines+markers',name:'Selected-feature estimate',line:{color:categoryColors['Driving Assistance'],width:3},marker:{color:categoryColors['Driving Assistance'],size:8},customdata:impact,hovertemplate:'%{x}: %{y:,.0f} vehicles<br>Feature uplift: %{customdata:.1f}%<extra></extra>',legendgroup:'selected'}
  ];
  const bmwLayout=layout(`${model} · ${source.label} · ${selected.length} selected feature${selected.length===1?'':'s'}`,colors);
  bmwLayout.xaxis={title:'Calendar year',tickmode:'array',tickvals:[2024,2025,2030,2031,2032,2033],gridcolor:colors.line,zeroline:false};
  bmwLayout.yaxis={title:'Estimated annual sales (vehicles)',rangemode:'tozero',gridcolor:colors.line,zeroline:false,tickformat:','};
  Plotly.react('bmw-sales',traces,bmwLayout,{responsive:true,displaylogo:false});
  bmwStatus.textContent=`${model} · ${source.label}: ${(cagr*100).toFixed(1)}% baseline CAGR · ${featureLabel}`;
}
function render() {
  const colors=theme();
  const categoryTraces=[];
  for (const row of categoryRows().filter(row=>visibleCategories.has(row.category))) {
    categoryTraces.push(...bandTraces(row.category,row.estimate,row.lower,row.upper,categoryColors[row.category],'solid',`${row.features} features · up to ${row.pairs} source pairs`));
  }
  const featureTraces=[];
  const selected=rows.filter(row=>visibleCategories.has(row.category)&&visibleFeatures.has(row.feature));
  selected.forEach((row,index)=>featureTraces.push(...bandTraces(row.feature,row.estimate,row.lower,row.upper,categoryColors[row.category],dashes[index%dashes.length],`${row.category} · ${row.pairs} source pairs`)));
  Plotly.react('categories',categoryTraces,layout('Category impact over four years',colors),{responsive:true,displaylogo:false});
  Plotly.react('features',featureTraces,layout('Feature impact over four years',colors),{responsive:true,displaylogo:false});
  renderBmw(colors);
  status.textContent=`${inputPairCount} input pairs analyzed · ${categoryTraces.length/3} visible categories · ${selected.length} visible features`;
}
bmwModel.addEventListener('change',()=>{updateMarkets();render();});
bmwMarket.addEventListener('change',render);
window.matchMedia('(prefers-color-scheme: dark)').addEventListener('change',render);
render();
</script>
</body>
</html>'''
    output.write_text(
        template.replace("__PAIR_COUNT__", str(pair_count))
        .replace("__ROW_COUNT__", str(row_count))
        .replace("__FEATURES__", payload).replace("__COLORS__", colors)
        .replace("__BMW_SALES__", bmw_sales),
        encoding="utf-8",
    )


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path, nargs="?", default=DEFAULT_INPUT)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--bmw-sales", type=Path, default=DEFAULT_BMW_SALES)
    args = parser.parse_args()
    records = read_versions(args.input)
    comparisons = version_sales_differences(records)
    snapshots = build_snapshots(args.input)
    bmw_sales_history = load_bmw_sales_history(args.bmw_sales)
    render_html(
        snapshots,
        args.output,
        input_pair_count=len(comparisons),
        input_row_count=len(records),
        bmw_sales_history=bmw_sales_history,
    )
    print(f"wrote three interactive Plotly uncertainty plots to {args.output}")


if __name__ == "__main__":
    main()
