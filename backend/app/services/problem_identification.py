"""Identify, visualize, and persist customer problems from the G70 feedback workbook."""

from __future__ import annotations

import json
import re
import textwrap
from collections import Counter, defaultdict
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openai import OpenAI
from openpyxl import load_workbook
from pydantic import BaseModel, Field

from app.paths import RUNTIME_DIR, get_openai_api_key, get_openai_model


OUTPUT_DIR = RUNTIME_DIR / "agent-problems"
OUTPUT_FILE = OUTPUT_DIR / "g70-problems.json"
FIGURE_FILE = OUTPUT_DIR / "g70-problems.png"
PROBLEM_FEEDBACK_TYPES = {"Difficult to Use", "Wants", "Defect"}
BROAD_LABELS = {
    "Body equipment, overall technology",
    "ConnectedDrive and Infotainment, overall technology",
    "Drivetrain",
    "Exterior",
    "General enquiry: customer feedback",
    "Seats",
}

CATEGORIES = (
    "User Problems",
    "Range & Charging",
    "Digital Experience",
    "Driving Experience",
    "Body Operation",
    "Cabin Experience",
)

FEATURES_BY_CATEGORY = {
    "User Problems": ["market_dynamics"],
    "Range & Charging": ["hv_800v", "ultra_charging", "v2g"],
    "Digital Experience": ["cockpit_llm", "phone_car_eco", "central_compute", "ota"],
    "Driving Experience": ["e2e_vla", "urban_noa", "active_suspension", "steer_by_wire", "lidar"],
    "Body Operation": ["smart_lighting"],
    "Cabin Experience": ["smart_seats"],
}

PROBLEM_COPY = {
    "Vehicle, climatization": ("Climate control usability and reliability", "Customers report difficulty operating climate controls, unmet climate needs, or functions that do not work as expected."),
    "Personal eSIM": ("Personal eSIM setup and connection failures", "Customers report difficult activation, unstable connections, or failures when using the vehicle's personal eSIM."),
    "Operating concept, operating system": ("Operating system interaction complexity", "Customers find the vehicle operating concept difficult to learn, navigate, or configure."),
    "Electric Range": ("Electric range confidence gap", "Customers report that available or communicated electric range does not meet their usage needs or expectations."),
    "Touch screen, operation": ("Touchscreen operation complexity", "Customers find touchscreen controls difficult to locate, understand, or operate reliably."),
    "Automatic engine start-stop function (MSA)": ("Automatic start-stop control friction", "Customers report unwanted start-stop behavior, difficult controls, or an unmet need for more predictable operation."),
    "Door, open / close": ("Door opening and closing friction", "Customers report difficulty, missing capability, or unreliable behavior when opening and closing doors."),
    "Panorama glass roof": ("Panoramic roof usability and reliability", "Customers report difficult controls, missing capability, or faults affecting the panoramic glass roof."),
    "Braking": ("Braking behavior and control concerns", "Customers report braking behavior that is difficult to manage, does not meet expectations, or may malfunction."),
    "My BMW App / MINI App": ("My BMW App usability and reliability", "Customers report difficult workflows, missing capability, or failures in the companion app."),
    "My Modes / driving experience": ("Driving mode configuration friction", "Customers report difficulty configuring driving modes or unmet expectations for the resulting vehicle experience."),
    "Wheels / tires": ("Wheel and tire capability gap", "Customers report missing capability, faults, or usability concerns related to wheels and tires."),
    "Apple CarPlay": ("Apple CarPlay connection and usability", "Customers report difficult setup, unstable connections, or failures while using Apple CarPlay."),
    "Public charging / external providers": ("Public charging access and compatibility", "Customers report difficult public-charging workflows, provider compatibility gaps, or charging failures."),
}

POSSIBLE_FEATURE_OVERRIDES = {
    "Vehicle, climatization": ["smart_seats"],
    "Personal eSIM": ["phone_car_eco", "ota"],
    "Operating concept, operating system": ["cockpit_llm", "central_compute", "ota"],
    "Electric Range": ["hv_800v", "ultra_charging"],
    "Touch screen, operation": ["cockpit_llm", "central_compute"],
    "Automatic engine start-stop function (MSA)": ["cockpit_llm", "central_compute"],
    "Door, open / close": ["smart_lighting"],
    "Panorama glass roof": ["smart_seats"],
    "Braking": ["steer_by_wire"],
    "My BMW App / MINI App": ["phone_car_eco", "ota"],
    "My Modes / driving experience": ["cockpit_llm", "active_suspension"],
    "Wheels / tires": ["active_suspension", "steer_by_wire"],
}


class ProblemEvidence(BaseModel):
    feedback_id: str
    feedback_type: str
    quote: str


class IdentifiedProblem(BaseModel):
    id: str
    name: str
    category: str
    problem_statement: str
    mention_count: int = Field(ge=1)
    severity: str
    affected_feature_ids: list[str]
    possible_feature_ids: list[str]
    evidence: list[ProblemEvidence]


class ProblemList(BaseModel):
    problems: list[IdentifiedProblem]


def _slug(value: str) -> str:
    slug = re.sub(r"[^a-z0-9]+", "-", value.lower()).strip("-")
    return slug[:60] or "problem"


def _category_for(label: str) -> str:
    text = label.lower()
    rules = [
        ("Range & Charging", ("range", "charg", "battery", "electric", "energy", "consumption")),
        ("Driving Experience", ("assistant", "parking", "blind spot", "cruise", "collision", "lane", "camera", "drivetrain", "handling", "riding", "chassis", "brak", "steering", "suspension", "tire", "start-stop", "my modes")),
        ("Digital Experience", ("touch screen", "display", "operating", "voice", "navigation", "head-up", "gesture", "connecteddrive", "esim", "app", "bluetooth", "phone", "carplay", "connection", "software")),
        ("Cabin Experience", ("seat", "interior", "comfort", "climat", "massage", "audio", "sound", "storage", "roof", "blower")),
        ("Body Operation", ("door", "tailgate", "exterior", "light", "body equipment", "mirror", "wiper", "key")),
    ]
    for category, needles in rules:
        if any(needle in text for needle in needles):
            return category
    return "User Problems"


def _feature_ids(label: str, category: str) -> list[str]:
    text = label.lower()
    specific = [
        ("smart_seats", ("seat", "massage", "interior comfort")),
        ("smart_lighting", ("door", "tailgate", "light", "headlight")),
        ("active_suspension", ("riding", "suspension", "damping", "cushion")),
        ("steer_by_wire", ("steering", "braking", "chassis", "handling")),
        ("cockpit_llm", ("voice", "assistant", "operating", "touch screen", "display")),
        ("phone_car_eco", ("esim", "app", "bluetooth", "phone", "connecteddrive")),
        ("ota", ("software", "update", "ota")),
        ("e2e_vla", ("highway assistant", "driving assistant", "cruise")),
        ("urban_noa", ("parking", "city", "urban")),
        ("lidar", ("sensor", "camera", "blind spot")),
        ("hv_800v", ("electric range", "battery")),
        ("ultra_charging", ("charg",)),
    ]
    matches = [feature_id for feature_id, needles in specific if any(needle in text for needle in needles)]
    return matches[:3]


def _possible_feature_ids(label: str, category: str) -> list[str]:
    if label in POSSIBLE_FEATURE_OVERRIDES:
        return POSSIBLE_FEATURE_OVERRIDES[label]
    exact = _feature_ids(label, category)
    candidates = exact + FEATURES_BY_CATEGORY[category]
    return list(dict.fromkeys(candidates))[:3]


def _severity(types: Counter[str]) -> str:
    if types["Defect"]:
        return "High"
    if types["Difficult to Use"] >= types["Wants"]:
        return "Medium"
    return "Opportunity"


def _read_candidates(source: Path | BytesIO, limit: int = 24) -> list[dict]:
    if isinstance(source, Path) and not source.exists():
        raise FileNotFoundError(f"Feedback workbook not found: {source}")
    workbook = load_workbook(source, read_only=True, data_only=True)
    if "Feedback_Explorer" not in workbook.sheetnames:
        raise ValueError("Workbook must contain a Feedback_Explorer sheet")
    worksheet = workbook["Feedback_Explorer"]
    grouped: dict[str, dict] = defaultdict(lambda: {"types": Counter(), "evidence": []})
    for row in worksheet.iter_rows(min_row=2, values_only=True):
        feedback_type = str(row[7] or "").strip()
        if feedback_type not in PROBLEM_FEEDBACK_TYPES:
            continue
        label = str(row[10] or "").strip()
        if not label or label == "no_class_found" or label in BROAD_LABELS:
            continue
        bucket = grouped[label]
        bucket["types"][feedback_type] += 1
        if len(bucket["evidence"]) < 3:
            bucket["evidence"].append(
                {
                    "feedback_id": str(row[0] or ""),
                    "feedback_type": feedback_type,
                    "quote": re.sub(r"\s+", " ", str(row[8] or "")).strip()[:420],
                }
            )
    ranked = sorted(grouped.items(), key=lambda item: (-sum(item[1]["types"].values()), item[0]))[:limit]
    return [
        {
            "label": label,
            "mention_count": sum(payload["types"].values()),
            "feedback_types": dict(payload["types"]),
            "evidence": payload["evidence"],
        }
        for label, payload in ranked
    ]


PROBLEM_SYSTEM_PROMPT = """You are a BMW product problem-identification agent. Convert aggregated G70 customer feedback into a concise, traceable problem list.

Rules:
- Return JSON only with a top-level `problems` array.
- Describe customer problems, not solutions, features, vague themes, or marketing language.
- Preserve the supplied mention_count and evidence objects exactly.
- Use exactly one of the supplied six customer-problem categories.
- Use only supplied Feature Radar IDs in affected_feature_ids.
- possible_feature_ids may include up to three plausible Feature Radar candidates, but must use only supplied IDs.
- Do not invent facts, affected vehicles, evidence, counts, or availability claims.
- Merge obvious synonyms only when the supplied evidence supports the merge.
- Return the 12 most important problems, prioritizing defects, frequency, and usability friction.
"""


def _llm_problems(candidates: list[dict]) -> list[IdentifiedProblem]:
    allowed_features = sorted({item for values in FEATURES_BY_CATEGORY.values() for item in values})
    payload = {
        "vehicle": "BMW 7 Series / i7 (G70)",
        "categories": CATEGORIES,
        "feature_radar_ids": allowed_features,
        "candidate_problem_clusters": candidates,
    }
    response = OpenAI(api_key=get_openai_api_key()).chat.completions.create(
        model=get_openai_model(),
        response_format={"type": "json_object"},
        temperature=0.0,
        seed=42,
        messages=[
            {"role": "system", "content": PROBLEM_SYSTEM_PROMPT},
            {"role": "user", "content": json.dumps(payload, ensure_ascii=False)},
        ],
    )
    parsed = ProblemList.model_validate_json(response.choices[0].message.content or '{"problems": []}')
    return parsed.problems


def _deterministic_problems(candidates: list[dict]) -> list[IdentifiedProblem]:
    problems = []
    for candidate in candidates[:12]:
        label = candidate["label"]
        category = _category_for(label)
        types = Counter(candidate["feedback_types"])
        dominant = types.most_common(1)[0][0]
        fallback_wording = {
            "Defect": f"Customers report that {label.lower()} can malfunction or fail to work as expected.",
            "Difficult to Use": f"Customers find {label.lower()} difficult or unintuitive to use.",
            "Wants": f"Customers report an unmet need related to {label.lower()}.",
        }[dominant]
        name, wording = PROBLEM_COPY.get(label, (f"{label} customer friction", fallback_wording))
        problems.append(
            IdentifiedProblem(
                id=f"g70-{_slug(label)}",
                name=name,
                category=category,
                problem_statement=wording,
                mention_count=candidate["mention_count"],
                severity=_severity(types),
                affected_feature_ids=_feature_ids(label, category),
                possible_feature_ids=_possible_feature_ids(label, category),
                evidence=[ProblemEvidence.model_validate(item) for item in candidate["evidence"]],
            )
        )
    return problems


def _save_figure(problems: list[IdentifiedProblem]) -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    shown = problems[:12]
    fig, ax = plt.subplots(figsize=(15.5, 8.2))
    fig.patch.set_facecolor("#0b1020")
    ax.set_facecolor("#0b1020")
    names = [textwrap.fill(problem.name.replace(" customer friction", ""), width=19) for problem in shown]
    counts = [problem.mention_count for problem in shown]
    maximum = max(counts, default=1)
    minimum = min(counts, default=0)
    spread = max(maximum - minimum, 1)
    colors = plt.cm.Blues([0.3 + 0.65 * ((count - minimum) / spread) for count in counts])
    bars = ax.bar(range(len(shown)), counts, color=colors, width=0.72)
    ax.bar_label(bars, padding=5, color="#e2e8f0", fontsize=10, fontweight="bold")
    ax.set_title("BMW G70 customer problems", loc="left", color="white", fontsize=20, fontweight="bold", pad=18)
    ax.set_ylabel("Problem mentions", color="#94a3b8", labelpad=12)
    ax.set_xticks(range(len(shown)), names, rotation=34, ha="right")
    ax.tick_params(colors="#cbd5e1", labelsize=10)
    for spine in ax.spines.values():
        spine.set_visible(False)
    ax.grid(axis="y", color="#334155", alpha=0.45, linewidth=0.8)
    ax.set_axisbelow(True)
    fig.tight_layout()
    fig.savefig(FIGURE_FILE, dpi=180, bbox_inches="tight", facecolor=fig.get_facecolor())
    plt.close(fig)


def identify_uploaded_problems(filename: str, content: bytes) -> dict:
    """Analyze an uploaded workbook, then save its structured JSON and PNG outputs."""
    if not filename.lower().endswith(".xlsx"):
        raise ValueError("Problem identification accepts .xlsx workbooks only")
    if not content:
        raise ValueError("The uploaded workbook is empty")
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    if "Feedback_Explorer" not in workbook.sheetnames:
        raise ValueError("Workbook must contain a Feedback_Explorer sheet")
    feedback_rows = max(workbook["Feedback_Explorer"].max_row - 1, 0)
    candidates = _read_candidates(BytesIO(content))
    api_key = get_openai_api_key()
    problems = _llm_problems(candidates) if api_key else _deterministic_problems(candidates)
    _save_figure(problems)
    output = {
        "vehicle": "G70",
        "source_file": Path(filename).name,
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "generation_mode": f"OpenAI · {get_openai_model()}" if api_key else "Deterministic fallback",
        "problem_categories": list(CATEGORIES),
        "feedback_rows": feedback_rows,
        "problem_rows": sum(candidate["mention_count"] for candidate in candidates),
        "figure_url": "/api/artifacts/agent-problems/g70-problems.png",
        "problems": [problem.model_dump() for problem in problems],
    }
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    OUTPUT_FILE.write_text(json.dumps(output, ensure_ascii=False, indent=2), encoding="utf-8")
    return output


def get_saved_g70_problems() -> dict:
    if not OUTPUT_FILE.exists():
        raise FileNotFoundError("No saved problem analysis yet. Upload a workbook in the AI Agent tab first.")
    return json.loads(OUTPUT_FILE.read_text(encoding="utf-8"))
